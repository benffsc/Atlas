#!/usr/bin/env python3
"""
Ingest ClinicHQ upcoming appointments XLSX into trapper.clinichq_upcoming_appointments.

REL_012: Windowed snapshot semantics
- Each import generates a run_id
- Coverage window is parsed from filename or data
- Only appointments WITHIN the coverage window that weren't seen are marked stale
- Appointments OUTSIDE the coverage window are untouched

Usage:
    python ingest_clinichq_upcoming_appointments.py --file <xlsx_path> [--schema trapper] [--dry-run] [--verbose]
"""
import argparse
import hashlib
import os
import re
import sys
import uuid
from dataclasses import dataclass
from datetime import datetime, date
from pathlib import Path
from typing import Any, Optional, Tuple

import psycopg
from dotenv import load_dotenv

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl required. Install with: pip install openpyxl", file=sys.stderr)
    sys.exit(2)


def norm_ws(s: str) -> str:
    if s is None:
        return ""
    return re.sub(r"\s+", " ", str(s)).strip()


def parse_date(val) -> Optional[str]:
    """Parse date from various formats."""
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.date().isoformat()
    if isinstance(val, date):
        return val.isoformat()
    s = str(val).strip()
    if not s:
        return None
    for fmt in ["%m/%d/%Y", "%Y-%m-%d", "%d/%m/%Y"]:
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def parse_int(val) -> Optional[int]:
    if val is None:
        return None
    try:
        return int(val)
    except (ValueError, TypeError):
        return None


def normalize_phone(val) -> Optional[str]:
    if val is None:
        return None
    digits = re.sub(r"\D+", "", str(val))
    if len(digits) == 11 and digits.startswith("1"):
        digits = digits[1:]
    return digits if digits else None


def compute_source_pk(row: dict) -> str:
    """Compute a stable source primary key for upserts.

    MEGA_008: source_pk is the stable identifier for a record.
    Uses appt_number if available (and non-zero), otherwise computed hash.
    """
    appt_number = row.get("appt_number")
    if appt_number is not None and appt_number > 0:
        return str(appt_number)

    # Fallback: hash of stable identifying fields
    appt_date = str(row.get("appt_date", ""))
    client_name = f"{norm_ws(row.get('client_first_name', ''))} {norm_ws(row.get('client_last_name', ''))}".lower()
    address = norm_ws(row.get("client_address", "")).lower()
    animal = norm_ws(row.get("animal_name", "")).lower()

    key_string = f"{appt_date}|{client_name}|{address}|{animal}"
    return f"hash:{hashlib.sha256(key_string.encode()).hexdigest()[:32]}"


def compute_row_hash(row: dict) -> str:
    """Compute a hash of the row content for change detection.

    MEGA_008: row_hash is for detecting changes, NOT for uniqueness.
    When row_hash differs but source_pk is the same, we update the existing record.
    """
    # Hash all significant content fields
    content_fields = [
        str(row.get("appt_date", "")),
        str(row.get("appt_number", "")),
        norm_ws(row.get("client_first_name", "")),
        norm_ws(row.get("client_last_name", "")),
        norm_ws(row.get("client_address", "")),
        norm_ws(row.get("animal_name", "")),
        norm_ws(row.get("client_email", "")),
        str(row.get("client_cell_phone", "")),
        str(row.get("client_phone", "")),
    ]
    content_string = "|".join(content_fields)
    return hashlib.sha256(content_string.encode()).hexdigest()[:32]


def parse_coverage_window_from_filename(filename: str) -> Tuple[Optional[date], Optional[date]]:
    """Extract date range from filename like 'clinichq_appts_2025-08-01_2026-02-28__pending.xlsx'.

    REL_012: Parse YYYY-MM-DD patterns from filename to determine coverage window.
    """
    # Look for patterns like YYYY-MM-DD_to_YYYY-MM-DD or YYYY-MM-DD_YYYY-MM-DD
    date_pattern = r'(\d{4}-\d{2}-\d{2})'
    matches = re.findall(date_pattern, filename)

    if len(matches) >= 2:
        try:
            start = datetime.strptime(matches[0], "%Y-%m-%d").date()
            end = datetime.strptime(matches[1], "%Y-%m-%d").date()
            return (start, end)
        except ValueError:
            pass

    return (None, None)


@dataclass
class Counters:
    rows_processed: int = 0
    rows_inserted: int = 0
    rows_updated: int = 0
    rows_skipped_blank: int = 0
    rows_staled: int = 0


def read_xlsx(file_path: Path) -> list[dict]:
    """Read XLSX file and return list of row dicts."""
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active

    rows_iter = ws.iter_rows()
    header_row = next(rows_iter)
    headers = [cell.value for cell in header_row]

    data = []
    for row in rows_iter:
        values = [cell.value for cell in row]
        row_dict = dict(zip(headers, values))
        data.append(row_dict)

    wb.close()
    return data


def is_blank_row(row: dict) -> bool:
    return all(v is None or str(v).strip() == "" for v in row.values())


def check_columns_exist(conn, schema: str) -> bool:
    """Check if windowed snapshot columns exist (MIG_263 applied)."""
    with conn.cursor() as cur:
        cur.execute("""
            SELECT COUNT(*) FROM information_schema.columns
            WHERE table_schema = %s
            AND table_name = 'clinichq_upcoming_appointments'
            AND column_name IN ('last_seen_run_id', 'is_current', 'stale_at')
        """, (schema,))
        result = cur.fetchone()
        return result[0] == 3


def check_ingest_runs_table(conn, schema: str) -> bool:
    """Check if ingest_runs table exists (MIG_263 applied)."""
    with conn.cursor() as cur:
        cur.execute("""
            SELECT EXISTS (
                SELECT 1 FROM information_schema.tables
                WHERE table_schema = %s AND table_name = 'ingest_runs'
            )
        """, (schema,))
        result = cur.fetchone()
        return result[0]


def upsert_appointment(
    conn,
    schema: str,
    row: dict,
    source_file: str,
    run_id: str,
    has_windowed_columns: bool,
    dry_run: bool,
    verbose: bool
) -> tuple[bool, bool]:
    """Upsert a single appointment. Returns (inserted, updated).

    MEGA_008: Uses source_pk (stable identifier) for ON CONFLICT, not row_hash.
    REL_012: Also sets last_seen_run_id and is_current when columns exist.
    """
    # Map XLSX columns to our fields
    appt_date = parse_date(row.get("Date"))
    appt_number = parse_int(row.get("Number"))
    animal_name = norm_ws(row.get("Animal Name"))
    ownership_type = norm_ws(row.get("Ownership"))
    client_type = norm_ws(row.get("ClientType"))
    client_first_name = norm_ws(row.get("Owner First Name"))
    client_last_name = norm_ws(row.get("Owner Last Name"))
    client_address = norm_ws(row.get("Owner Address"))
    client_cell_phone = normalize_phone(row.get("Owner Cell Phone"))
    client_phone = normalize_phone(row.get("Owner Phone"))
    client_email = norm_ws(row.get("Owner Email"))

    # Build normalized row for hashing
    normalized = {
        "appt_date": appt_date,
        "appt_number": appt_number,
        "client_first_name": client_first_name,
        "client_last_name": client_last_name,
        "client_address": client_address,
        "animal_name": animal_name,
        "client_email": client_email,
        "client_cell_phone": client_cell_phone,
        "client_phone": client_phone,
    }
    source_pk = compute_source_pk(normalized)
    row_hash = compute_row_hash(normalized)

    if verbose:
        print(f"  Source PK: {source_pk}")
        print(f"  Row hash: {row_hash}")
        print(f"  Date: {appt_date}, Client: {client_first_name} {client_last_name}, Animal: {animal_name}")

    if appt_date is None:
        if verbose:
            print("  Skipping: no appt_date")
        return False, False

    with conn.cursor() as cur:
        if has_windowed_columns:
            # REL_012: Include windowed snapshot columns
            sql = f"""
            INSERT INTO {schema}.clinichq_upcoming_appointments (
                source_file, source_row_hash, source_system, source_pk,
                appt_date, appt_number,
                client_first_name, client_last_name, client_address,
                client_cell_phone, client_phone, client_email, client_type,
                animal_name, ownership_type,
                first_seen_at, last_seen_at, last_seen_run_id, is_current, stale_at
            ) VALUES (
                %s, %s, %s, %s,
                %s, %s,
                %s, %s, %s,
                %s, %s, %s, %s,
                %s, %s,
                NOW(), NOW(), %s, true, NULL
            )
            ON CONFLICT (source_system, source_pk) DO UPDATE SET
                source_file = EXCLUDED.source_file,
                source_row_hash = EXCLUDED.source_row_hash,
                appt_date = EXCLUDED.appt_date,
                appt_number = COALESCE(EXCLUDED.appt_number, {schema}.clinichq_upcoming_appointments.appt_number),
                client_first_name = COALESCE(EXCLUDED.client_first_name, {schema}.clinichq_upcoming_appointments.client_first_name),
                client_last_name = COALESCE(EXCLUDED.client_last_name, {schema}.clinichq_upcoming_appointments.client_last_name),
                client_address = COALESCE(EXCLUDED.client_address, {schema}.clinichq_upcoming_appointments.client_address),
                client_cell_phone = COALESCE(EXCLUDED.client_cell_phone, {schema}.clinichq_upcoming_appointments.client_cell_phone),
                client_phone = COALESCE(EXCLUDED.client_phone, {schema}.clinichq_upcoming_appointments.client_phone),
                client_email = COALESCE(EXCLUDED.client_email, {schema}.clinichq_upcoming_appointments.client_email),
                animal_name = COALESCE(EXCLUDED.animal_name, {schema}.clinichq_upcoming_appointments.animal_name),
                last_seen_at = NOW(),
                last_seen_run_id = %s,
                is_current = true,
                stale_at = NULL,
                updated_at = NOW()
            RETURNING (xmax = 0) AS inserted;
            """
            cur.execute(sql, (
                source_file, row_hash, "clinichq", source_pk,
                appt_date, appt_number,
                client_first_name or None, client_last_name or None, client_address or None,
                client_cell_phone, client_phone, client_email or None, client_type or None,
                animal_name or None, ownership_type or None,
                run_id, run_id
            ))
        else:
            # Legacy mode: no windowed snapshot columns
            sql = f"""
            INSERT INTO {schema}.clinichq_upcoming_appointments (
                source_file, source_row_hash, source_system, source_pk,
                appt_date, appt_number,
                client_first_name, client_last_name, client_address,
                client_cell_phone, client_phone, client_email, client_type,
                animal_name, ownership_type,
                first_seen_at, last_seen_at
            ) VALUES (
                %s, %s, %s, %s,
                %s, %s,
                %s, %s, %s,
                %s, %s, %s, %s,
                %s, %s,
                NOW(), NOW()
            )
            ON CONFLICT (source_system, source_pk) DO UPDATE SET
                source_file = EXCLUDED.source_file,
                source_row_hash = EXCLUDED.source_row_hash,
                appt_date = EXCLUDED.appt_date,
                appt_number = COALESCE(EXCLUDED.appt_number, {schema}.clinichq_upcoming_appointments.appt_number),
                client_first_name = COALESCE(EXCLUDED.client_first_name, {schema}.clinichq_upcoming_appointments.client_first_name),
                client_last_name = COALESCE(EXCLUDED.client_last_name, {schema}.clinichq_upcoming_appointments.client_last_name),
                client_address = COALESCE(EXCLUDED.client_address, {schema}.clinichq_upcoming_appointments.client_address),
                client_cell_phone = COALESCE(EXCLUDED.client_cell_phone, {schema}.clinichq_upcoming_appointments.client_cell_phone),
                client_phone = COALESCE(EXCLUDED.client_phone, {schema}.clinichq_upcoming_appointments.client_phone),
                client_email = COALESCE(EXCLUDED.client_email, {schema}.clinichq_upcoming_appointments.client_email),
                animal_name = COALESCE(EXCLUDED.animal_name, {schema}.clinichq_upcoming_appointments.animal_name),
                last_seen_at = NOW(),
                updated_at = NOW()
            RETURNING (xmax = 0) AS inserted;
            """
            cur.execute(sql, (
                source_file, row_hash, "clinichq", source_pk,
                appt_date, appt_number,
                client_first_name or None, client_last_name or None, client_address or None,
                client_cell_phone, client_phone, client_email or None, client_type or None,
                animal_name or None, ownership_type or None
            ))

        result = cur.fetchone()
        inserted = bool(result[0]) if result else False
        return inserted, not inserted


def mark_stale_in_window(
    conn,
    schema: str,
    run_id: str,
    coverage_start: date,
    coverage_end: date,
    dry_run: bool,
    verbose: bool
) -> int:
    """Mark appointments as stale if they're in the coverage window but weren't seen in this run.

    REL_012: Windowed snapshot semantics - only expire within the import's date range.
    """
    with conn.cursor() as cur:
        sql = f"""
        UPDATE {schema}.clinichq_upcoming_appointments
        SET is_current = false, stale_at = NOW()
        WHERE is_current = true
          AND appt_date >= %s
          AND appt_date <= %s
          AND (last_seen_run_id IS NULL OR last_seen_run_id != %s)
        """
        cur.execute(sql, (coverage_start, coverage_end, run_id))
        staled_count = cur.rowcount

        if verbose:
            print(f"  Marked {staled_count} appointments as stale in window {coverage_start} to {coverage_end}")

        return staled_count


def log_ingest_run(
    conn,
    schema: str,
    run_id: str,
    source_file: str,
    coverage_start: Optional[date],
    coverage_end: Optional[date],
    counters: Counters,
    dry_run: bool
):
    """Log this ingest run to the ingest_runs table."""
    with conn.cursor() as cur:
        sql = f"""
        INSERT INTO {schema}.ingest_runs (
            run_id, source_type, source_file,
            started_at, completed_at,
            coverage_start, coverage_end,
            rows_processed, rows_inserted, rows_updated, rows_staled,
            notes
        ) VALUES (
            %s, %s, %s,
            NOW(), NOW(),
            %s, %s,
            %s, %s, %s, %s,
            %s
        )
        """
        notes = "dry-run" if dry_run else None
        cur.execute(sql, (
            run_id, "clinichq_upcoming", source_file,
            coverage_start, coverage_end,
            counters.rows_processed, counters.rows_inserted, counters.rows_updated, counters.rows_staled,
            notes
        ))


def main():
    load_dotenv()

    ap = argparse.ArgumentParser(description="Ingest ClinicHQ upcoming appointments XLSX")
    ap.add_argument("--file", required=True, help="Path to XLSX file")
    ap.add_argument("--schema", default="trapper", help="DB schema (default: trapper)")
    ap.add_argument("--dry-run", action="store_true", help="Print what would happen, no DB writes")
    ap.add_argument("--verbose", action="store_true", help="Print detailed output")
    args = ap.parse_args()

    db_url = os.getenv("DATABASE_URL", "").strip()
    if not db_url:
        print("ERROR: DATABASE_URL not set in .env", file=sys.stderr)
        sys.exit(2)

    xlsx_path = Path(args.file)
    if not xlsx_path.exists():
        print(f"ERROR: File not found: {xlsx_path}", file=sys.stderr)
        sys.exit(2)

    source_file = xlsx_path.name
    counters = Counters()

    # REL_012: Generate unique run_id for this import
    run_id = str(uuid.uuid4())

    if args.dry_run:
        print(f"[DRY RUN] Would ingest: {xlsx_path}")

    # REL_012: Parse coverage window from filename
    coverage_start, coverage_end = parse_coverage_window_from_filename(source_file)
    if coverage_start and coverage_end:
        print(f"Coverage window from filename: {coverage_start} to {coverage_end}")

    # Read XLSX
    print(f"Reading {xlsx_path}...")
    rows = read_xlsx(xlsx_path)
    print(f"Found {len(rows)} rows")
    print(f"Run ID: {run_id}")

    with psycopg.connect(db_url) as conn:
        conn.execute("SELECT 1")  # smoke test

        # Check if MIG_263 columns exist
        has_windowed_columns = check_columns_exist(conn, args.schema)
        has_ingest_runs = check_ingest_runs_table(conn, args.schema)

        if has_windowed_columns:
            print("Windowed snapshot columns detected (MIG_263 applied)")
        else:
            print("NOTE: MIG_263 not applied. Running in legacy mode (no windowed snapshot).")

        # Track min/max dates seen in the data (fallback for coverage window)
        data_min_date = None
        data_max_date = None

        for row in rows:
            if is_blank_row(row):
                counters.rows_skipped_blank += 1
                continue

            counters.rows_processed += 1

            if args.verbose:
                print(f"\nProcessing row {counters.rows_processed}...")

            inserted, updated = upsert_appointment(
                conn, args.schema, row, source_file, run_id, has_windowed_columns, args.dry_run, args.verbose
            )

            if inserted:
                counters.rows_inserted += 1
            elif updated:
                counters.rows_updated += 1

            # Track data date range for coverage window fallback
            row_date_str = parse_date(row.get("Date"))
            if row_date_str:
                row_date = datetime.strptime(row_date_str, "%Y-%m-%d").date()
                if data_min_date is None or row_date < data_min_date:
                    data_min_date = row_date
                if data_max_date is None or row_date > data_max_date:
                    data_max_date = row_date

        # REL_012: Determine final coverage window
        if coverage_start is None:
            coverage_start = data_min_date
        if coverage_end is None:
            coverage_end = data_max_date

        if coverage_start and coverage_end:
            print(f"Final coverage window: {coverage_start} to {coverage_end}")

        # REL_012: Mark stale appointments within the coverage window
        if has_windowed_columns and coverage_start and coverage_end:
            staled = mark_stale_in_window(
                conn, args.schema, run_id, coverage_start, coverage_end, args.dry_run, args.verbose
            )
            counters.rows_staled = staled

        # Log the ingest run
        if has_ingest_runs and not args.dry_run:
            log_ingest_run(
                conn, args.schema, run_id, source_file,
                coverage_start, coverage_end, counters, args.dry_run
            )

        if args.dry_run:
            conn.rollback()
            print("\n[ROLLBACK - dry-run mode, no changes persisted]")
        else:
            conn.commit()

    print(f"\n=== ClinicHQ Upcoming Appointments Import Summary ===")
    print(f"File: {source_file}")
    print(f"Run ID: {run_id}")
    if coverage_start and coverage_end:
        print(f"Coverage: {coverage_start} to {coverage_end}")
    print(f"Rows processed:    {counters.rows_processed}")
    print(f"Rows inserted:     {counters.rows_inserted}")
    print(f"Rows updated:      {counters.rows_updated}")
    print(f"Rows staled:       {counters.rows_staled}")
    print(f"Rows skipped:      {counters.rows_skipped_blank}")
    if args.dry_run:
        print("[DRY RUN - no changes made]")
    print("=" * 52)


if __name__ == "__main__":
    main()
